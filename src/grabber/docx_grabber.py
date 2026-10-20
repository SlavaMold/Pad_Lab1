from pathlib import Path
from docx import Document as DocxDocument

from .errors import SourceUnavailableError, SourceNotFoundError, ParsingError
from .logger import get_logger
from .models import RawDocument
from .registry import Registry

logger = get_logger()

RAW_DOCX_DIR = Path("data/raw/docx")
RAW_DOCX_DIR.mkdir(parents=True, exist_ok=True)

RECIPE_TITLE_STYLES = {"Heading 1", "Заголовок 1", "Title", "Название"}
SUBSECTION_STYLES = {"Heading 2", "Heading 3", "Заголовок 2", "Заголовок 3"}


def read_docx_file(file_path: str) -> DocxDocument:
    """Открывает .docx файл. Поднимает исключения при отсутствии/повреждении файла."""
    path = Path(file_path)

    if not path.exists():
        raise SourceNotFoundError(f"Файл не найден: {file_path}")

    try:
        return DocxDocument(str(path))
    except Exception as e:
        # python-docx падает с общим Exception на повреждённых/неверных файлах
        raise ParsingError(f"Не удалось открыть {file_path}: {e}")


def split_into_recipes(doc: DocxDocument, file_path: str) -> list[tuple[str, str]]:
    recipes = []
    current_title = None
    current_lines = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        style_name = para.style.name

        if style_name in RECIPE_TITLE_STYLES:
            # Новый рецепт верхнего уровня — сохраняем предыдущий
            if current_title is not None and current_lines:
                recipes.append((current_title, "\n".join(current_lines)))
            current_title = text
            current_lines = []

        elif style_name in SUBSECTION_STYLES:
            # Подраздел (Ingredients/Directions и т.п.) — НЕ новый рецепт,
            # просто помечаем его текстовым маркером внутри текущего рецепта
            current_lines.append(f"\n## {text}")

        else:
            current_lines.append(text)

    if current_title is not None and current_lines:
        recipes.append((current_title, "\n".join(current_lines)))

    if not recipes:
        # Заголовков верхнего уровня не было вообще (как в recipe_2.docx) —
        # весь файл = один рецепт, а подразделы всё равно размечены через "## "
        all_lines = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            if para.style.name in SUBSECTION_STYLES:
                all_lines.append(f"\n## {text}")
            else:
                all_lines.append(text)

        all_text = "\n".join(all_lines)
        if len(all_text) < 50:
            raise ParsingError(f"Слишком мало текста в {file_path} ({len(all_text)} симв.)")

        fallback_title = Path(file_path).stem
        recipes.append((fallback_title, all_text))

    return recipes


def grab_docx_source(file_path: str, registry: Registry) -> list[RawDocument]:
    """
    Полный цикл для одного .docx файла:
    открыть -> разбить на рецепты -> проверить дубликаты -> сохранить.
    Возвращает список RawDocument для всех новых/изменённых рецептов из файла.
    """
    try:
        doc = read_docx_file(file_path)
        recipes = split_into_recipes(doc, file_path)
    except (SourceNotFoundError, ParsingError) as e:
        logger.error(f"Пропускаю файл {file_path}: {e}")
        return []

    saved_docs = []

    for title, content in recipes:
        # source_id должен быть уникальным для КАЖДОГО рецепта внутри файла,
        # иначе реестр перепутает разные рецепты одного файла между собой
        source_id = f"{file_path}::{title}"

        if not registry.is_new_or_changed(source_id=source_id, content=content):
            logger.info(f"Без изменений, пропускаю: {source_id}")
            continue

        safe_filename = registry.compute_hash(source_id)[:16] + ".txt"
        raw_path = RAW_DOCX_DIR / safe_filename

        with open(raw_path, "w", encoding="utf-8") as f:
            f.write(content)

        raw_doc = RawDocument(
            source_id=source_id,
            source_type="docx",
            title=title,
            content=content,
            content_hash=registry.compute_hash(content),
            raw_path=str(raw_path),
            extra_metadata={"source_file": file_path},
        )

        registry.update(source_id=source_id, content=content, raw_path=str(raw_path))
        logger.info(f"Сохранён новый/изменённый рецепт: {title} (из {file_path})")
        saved_docs.append(raw_doc)

    return saved_docs