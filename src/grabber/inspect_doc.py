import subprocess
from pathlib import Path
from docx import Document as DocxDocument


def convert_doc_to_docx(doc_path: str, output_dir: str = "data/temp_converted") -> str:
    """Конвертирует .doc в .docx через LibreOffice headless. Возвращает путь к .docx."""
    Path(output_dir).mkdir(parents=True, exist_ok=True)

    result = subprocess.run(
        ["soffice", "--headless", "--convert-to", "docx", "--outdir", output_dir, doc_path],
        capture_output=True, text=True, timeout=60
    )

    if result.returncode != 0:
        raise RuntimeError(f"Ошибка конвертации {doc_path}: {result.stderr}")

    converted_name = Path(doc_path).stem + ".docx"
    converted_path = Path(output_dir) / converted_name

    if not converted_path.exists():
        raise RuntimeError(f"Файл не появился после конвертации: {converted_path}")

    return str(converted_path)


def inspect_structure(docx_path: str):
    doc = DocxDocument(docx_path)

    print(f"\n{'='*60}\nФайл: {docx_path}\n{'='*60}")

    print(f"\n--- Параграфы (первые 30 непустых) ---")
    count = 0
    for p in doc.paragraphs:
        if p.text.strip():
            print(f"[{p.style.name}] {p.text.strip()[:60]}")
            count += 1
            if count >= 30:
                print("... (обрезано)")
                break

    print(f"\n--- Таблицы ---")
    print(f"Всего таблиц в документе: {len(doc.tables)}")
    for i, table in enumerate(doc.tables[:2]):  # смотрим только первые 2 таблицы
        print(f"\nТаблица {i}: {len(table.rows)} строк x {len(table.columns)} столбцов")
        for row_idx, row in enumerate(table.rows[:5]):  # первые 5 строк
            cells_text = [cell.text.strip()[:40] for cell in row.cells]
            print(f"  Строка {row_idx}: {cells_text}")


if __name__ == "__main__":
    # Подставь путь к ОДНОМУ реальному .doc файлу из твоей базы
    doc_path = "data/input_docs/meat1.doc"

    docx_path = convert_doc_to_docx(doc_path)
    inspect_structure(docx_path)